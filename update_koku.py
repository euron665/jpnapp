"""
KOKU Update Pipeline
──────────────────────────────────────────────────────────────────
Reads the Excel spreadsheets in this folder and rebuilds:
  • kanji.json, skgs.json, vocab.json, pos_tags.json
  • index.html  (inline JS constants)
  • sw.js       (bumps cache version so phone picks up changes)

Run: double-click update_koku.bat  (or  python update_koku.py)
"""

import openpyxl, json, re, os, sys, glob
from datetime import datetime
from zoneinfo import ZoneInfo

APP_DIR = os.path.dirname(os.path.abspath(__file__))

def find_latest(pattern):
    matches = glob.glob(os.path.join(APP_DIR, pattern))
    if not matches:
        print(f'ERROR: no file matching {pattern}')
        input('Press Enter to close...')
        sys.exit(1)
    latest = max(matches, key=os.path.getmtime)
    print(f'  Using: {os.path.basename(latest)}')
    return latest

VOCAB_XLSX  = find_latest('FINAL VOCAB LIST*.xlsx')
KANJI_XLSX  = find_latest('FINAL KANJI & SKG LIST*.xlsx')
POS_XLSX    = find_latest('FINAL POS LIST*.xlsx')
KATAKANA_XLSX = find_latest('FINAL KATAKANA VOCAB*.xlsx')
INDEX_HTML  = os.path.join(APP_DIR, 'index.html')
SW_JS       = os.path.join(APP_DIR, 'sw.js')

# ── helpers ───────────────────────────────────────────────────────────────────
def get_color(cell):
    """Return fgColor RGB string if explicitly set, else None."""
    try:
        fill = cell.fill
        if fill and fill.fgColor and fill.fgColor.type == 'rgb':
            rgb = fill.fgColor.rgb
            if rgb and rgb != '00000000':
                return rgb
    except Exception:
        pass
    return None

def parse_rm(val):
    """'reading\\nmeaning' → (reading, meaning)"""
    if not val: return '', ''
    parts = str(val).strip().split('\n')
    return (parts[0].strip(), parts[1].strip()) if len(parts) >= 2 else (parts[0].strip(), '')

def parse_vocab_cell(val):
    """'word\\nreading\\nmeaning' → dict or None"""
    if not val: return None
    parts = str(val).strip().split('\n')
    if len(parts) >= 3:
        return {'word': parts[0].strip(), 'reading': parts[1].strip(), 'meaning': parts[2].strip()}
    return None

def find_const_bounds(html, name):
    """
    Find start/end char positions of 'const NAME = <value>;'
    using bracket counting — safe for large JSON blobs.
    Returns (start, end) or (None, None) if not found.
    """
    m = re.search(r'const ' + re.escape(name) + r'\s*=\s*', html)
    if not m:
        return None, None

    start = m.start()
    i = m.end()
    depth = 0
    in_string = False
    string_char = None

    while i < len(html):
        c = html[i]
        if in_string:
            if c == '\\':
                i += 2
                continue
            if c == string_char:
                in_string = False
        else:
            if c in ('"', "'"):
                in_string = True
                string_char = c
            elif c in ('{', '['):
                depth += 1
            elif c in ('}', ']'):
                depth -= 1
                if depth == 0:
                    j = i + 1
                    while j < len(html) and html[j] in ' \t':
                        j += 1
                    end = j + 1 if j < len(html) and html[j] == ';' else i + 1
                    return start, end
        i += 1

    return None, None

def replace_const(html, name, value_str):
    """Safely replace a const declaration using bracket counting (no regex on data)."""
    start, end = find_const_bounds(html, name)
    if start is None:
        print(f"  WARNING: could not find const {name} in index.html")
        return html
    m = re.search(r'const ' + re.escape(name) + r'\s*=\s*', html[start:start + 200])
    prefix = html[start : start + m.end()]
    return html[:start] + prefix + value_str + ';' + html[end:]

# ── 1. KANJI + SKGS ───────────────────────────────────────────────────────────
print("Reading kanji spreadsheet...")
wb_k = openpyxl.load_workbook(KANJI_XLSX)
ws_k = wb_k['Sheet2']

kanji_list = []
skgs_dict  = {}

for r in range(1, ws_k.max_row + 1):
    kanji = ws_k.cell(r, 1).value
    rm    = ws_k.cell(r, 2).value
    skg   = ws_k.cell(r, 3).value
    if not kanji or not skg:
        continue
    reading, meaning = parse_rm(rm)
    kanji_list.append({'kanji': str(kanji), 'reading': reading, 'meaning': meaning, 'skg': str(skg)})
    skgs_dict.setdefault(str(skg), []).append(str(kanji))

skgs_list = [{'id': k, 'kanji': v} for k, v in sorted(skgs_dict.items())]
print(f"  {len(kanji_list)} kanji across {len(skgs_list)} SKGs")

# ── 2. VOCAB (color-based reading groups) ─────────────────────────────────────
print("Reading vocab spreadsheet...")
wb_v = openpyxl.load_workbook(VOCAB_XLSX)
ws_v = wb_v['Sheet1']

vocab_dict = {}
for r in range(1, ws_v.max_row + 1):
    kanji = ws_v.cell(r, 1).value
    if not kanji:
        continue

    # Col E = RG reading labels (colon-separated); vocab starts at col F
    rg_labels_raw = ws_v.cell(r, 5).value
    rg_labels = [s.strip() for s in str(rg_labels_raw).split(':')] if rg_labels_raw else []

    # Scan columns F onwards, group consecutive same-color cells into RGs
    rg_groups     = []   # list of lists of vocab items
    current_color = None
    current_group = []

    for c in range(6, ws_v.max_column + 1):
        cell = ws_v.cell(r, c)
        item = parse_vocab_cell(cell.value)
        if item is None:
            continue                       # skip empty cells
        color = get_color(cell)
        if color != current_color:
            if current_group:
                rg_groups.append(current_group)
            current_color = color
            current_group = [item]
        else:
            current_group.append(item)

    if current_group:
        rg_groups.append(current_group)

    rg_list = [
        {
            'rg_index':   i + 1,
            'rg_reading': rg_labels[i] if i < len(rg_labels) else items[0]['reading'][:3],
            'vocab':      items,
        }
        for i, items in enumerate(rg_groups)
    ]

    vocab_dict[str(kanji)] = {
        'kanji':    str(kanji),
        'rg_count': len(rg_list),
        'rgs':      rg_list,
    }

print(f"  {len(vocab_dict)} kanji vocab entries")

# ── 3. POS_TAGS ───────────────────────────────────────────────────────────────
print("Reading POS spreadsheet...")
POS_SHEET_MAP = {
    'い Adjectives':           ['i-adj'],
    'な Adjectives':           ['na-adj', 'noun'],
    'Ichidan Verbs':           ['ru-verb'],
    'Godan Verbs (except る)': ['u-verb'],
    'Godan Verbs (only る)':   ['u-verb-ru'],
    'Suru Verbs':              ['suru-verb'],
    'Adverbs':                 ['adverb'],
    'Expressions':             ['expression'],
}

wb_p     = openpyxl.load_workbook(POS_XLSX)
pos_tags = {}

for sheet_name, tags in POS_SHEET_MAP.items():
    if sheet_name not in wb_p.sheetnames:
        print(f"  WARNING: sheet '{sheet_name}' not found")
        continue
    ws_p  = wb_p[sheet_name]
    count = 0
    for r in range(1, ws_p.max_row + 1):
        word = ws_p.cell(r, 1).value
        if word:
            pos_tags[str(word).strip()] = tags
            count += 1
    print(f"  {sheet_name}: {count} words")

# Default untagged vocab words to noun
for data in vocab_dict.values():
    for rg in data['rgs']:
        for item in rg['vocab']:
            pos_tags.setdefault(item['word'], ['noun'])

print(f"  {len(pos_tags)} total POS entries")

# ── 3b. KATAKANA WORDS ────────────────────────────────────────────────────────
print("Reading katakana spreadsheet...")

_KB = {
 'ア':'a','イ':'i','ウ':'u','エ':'e','オ':'o','カ':'ka','キ':'ki','ク':'ku','ケ':'ke','コ':'ko',
 'ガ':'ga','ギ':'gi','グ':'gu','ゲ':'ge','ゴ':'go','サ':'sa','シ':'shi','ス':'su','セ':'se','ソ':'so',
 'ザ':'za','ジ':'ji','ズ':'zu','ゼ':'ze','ゾ':'zo','タ':'ta','チ':'chi','ツ':'tsu','テ':'te','ト':'to',
 'ダ':'da','ヂ':'ji','ヅ':'zu','デ':'de','ド':'do','ナ':'na','ニ':'ni','ヌ':'nu','ネ':'ne','ノ':'no',
 'ハ':'ha','ヒ':'hi','フ':'fu','ヘ':'he','ホ':'ho','バ':'ba','ビ':'bi','ブ':'bu','ベ':'be','ボ':'bo',
 'パ':'pa','ピ':'pi','プ':'pu','ペ':'pe','ポ':'po','マ':'ma','ミ':'mi','ム':'mu','メ':'me','モ':'mo',
 'ヤ':'ya','ユ':'yu','ヨ':'yo','ラ':'ra','リ':'ri','ル':'ru','レ':'re','ロ':'ro',
 'ワ':'wa','ヲ':'wo','ン':'n','ヴ':'vu',
}
_SY = {'ャ':'ya','ュ':'yu','ョ':'yo'}
_SV = {'ァ':'a','ィ':'i','ゥ':'u','ェ':'e','ォ':'o'}
_MAC = {'a':'ā','i':'ī','u':'ū','e':'ē','o':'ō'}

def _stem(b):
    if b == 'shi': return 'sh'
    if b == 'chi': return 'ch'
    if b == 'ji':  return 'j'
    if b == 'tsu': return 'ts'
    if b == 'fu':  return 'f'
    if b == 'u':   return 'w'
    if b == 'vu':  return 'v'
    return b[:-1] if b and b[-1] in 'aiueo' else b

def kata_to_romaji(word):
    ch = list(word); toks = []; i = 0
    while i < len(ch):
        c = ch[i]; n = ch[i+1] if i+1 < len(ch) else ''
        if c in _KB and n in _SY:
            b = _KB[c]
            if b in ('shi', 'chi', 'ji'):
                toks.append({'shi':'sh','chi':'ch','ji':'j'}[b] + {'ャ':'a','ュ':'u','ョ':'o'}[n])
            else:
                toks.append(b[:-1] + _SY[n])
            i += 2; continue
        if c in _KB and n in _SV:
            toks.append(_stem(_KB[c]) + _SV[n]); i += 2; continue
        if c == 'ー': toks.append('ー'); i += 1; continue
        if c == 'ッ': toks.append('ッ'); i += 1; continue
        toks.append(_KB.get(c, c)); i += 1
    out = ''; dbl = False
    for t in toks:
        if t == 'ッ': dbl = True; continue
        if t == 'ー':
            if out and out[-1] in 'aiueo': out = out[:-1] + _MAC[out[-1]]
            continue
        if dbl and t: out += t[0]; dbl = False
        out += t
    return out

# accept-answer synonyms: stored meaning -> pipe-separated acceptable answers
KATA_SYN = {
 "television":"television|tv", "PC; computer":"pc|computer|personal computer",
 "e-mail":"email|e-mail|mail", "air conditioner":"air conditioner|aircon|ac",
 "convenience store":"convenience store|conbini|konbini",
 "department store":"department store|depato",
 "smartphone":"smartphone|smart phone|phone", "remote control":"remote control|remote",
 "credit card":"credit card|card",
}

wb_kt = openpyxl.load_workbook(KATAKANA_XLSX)
ws_kt = wb_kt['Katakana'] if 'Katakana' in wb_kt.sheetnames else wb_kt.active
katakana = []
for r in range(2, ws_kt.max_row + 1):          # skip header row 1
    w = ws_kt.cell(r, 1).value
    if not w:
        continue
    m    = str(ws_kt.cell(r, 2).value or '').strip()
    cat  = str(ws_kt.cell(r, 3).value or '').strip()
    note = str(ws_kt.cell(r, 4).value or '').strip()
    katakana.append({
        'word':     str(w).strip(),
        'meaning':  KATA_SYN.get(m, m),
        'category': cat,
        'note':     note,
        'romaji':   kata_to_romaji(str(w).strip()),
    })
print(f"  {len(katakana)} katakana words")

# ── 3c. SAME-MEANING (different kanji) ────────────────────────────────────────
print("Reading same-meaning list...")
sm_matches = glob.glob(os.path.join(APP_DIR, 'FINAL SAME MEANING LIST*.xlsx'))
same_meaning = []
if sm_matches:
    sm_file = max(sm_matches, key=os.path.getmtime)
    print(f'  Using: {os.path.basename(sm_file)}')
    wb_sm = openpyxl.load_workbook(sm_file)
    ws_sm = wb_sm[wb_sm.sheetnames[0]]
    for r in range(2, ws_sm.max_row + 1):           # skip header row 1
        meaning = ws_sm.cell(r, 1).value
        if not meaning:
            continue
        words = []
        for i in range(6):                          # 6 word slots: cols 2-13
            w  = ws_sm.cell(r, 2 + i * 2).value
            rd = ws_sm.cell(r, 3 + i * 2).value
            if w:
                words.append({'w': str(w).strip(), 'r': str(rd).strip() if rd else ''})
        if len(words) < 2:
            continue
        note = ws_sm.cell(r, 14).value              # Note lives in col 14
        same_meaning.append({'m': str(meaning).strip(), 'words': words,
                             'note': str(note).strip() if note else ''})
    print(f"  {len(same_meaning)} same-meaning groups")
else:
    print("  (no same-meaning list found — leaving SAME_MEANING empty)")

# ── 4. Write JSON files ───────────────────────────────────────────────────────
print("Writing JSON files...")
with open(os.path.join(APP_DIR, 'kanji.json'),    'w', encoding='utf-8') as f:
    json.dump(kanji_list, f, ensure_ascii=False, indent=2)
with open(os.path.join(APP_DIR, 'skgs.json'),     'w', encoding='utf-8') as f:
    json.dump(skgs_list,  f, ensure_ascii=False, indent=2)
with open(os.path.join(APP_DIR, 'vocab.json'),    'w', encoding='utf-8') as f:
    json.dump(vocab_dict, f, ensure_ascii=False, indent=2)
with open(os.path.join(APP_DIR, 'pos_tags.json'), 'w', encoding='utf-8') as f:
    json.dump(pos_tags,   f, ensure_ascii=False, indent=2)
with open(os.path.join(APP_DIR, 'katakana.json'), 'w', encoding='utf-8') as f:
    json.dump(katakana,   f, ensure_ascii=False, indent=2)
with open(os.path.join(APP_DIR, 'same_meaning.json'), 'w', encoding='utf-8') as f:
    json.dump(same_meaning, f, ensure_ascii=False, indent=2)
print("  Done")

# ── 5. Rebuild index.html inline constants ────────────────────────────────────
print("Rebuilding index.html...")
with open(INDEX_HTML, 'r', encoding='utf-8') as f:
    html = f.read()

kanji_js    = json.dumps(kanji_list, ensure_ascii=False, separators=(', ', ': '))
skgs_js     = json.dumps(skgs_list,  ensure_ascii=False, separators=(', ', ': '))
vocab_js    = json.dumps(vocab_dict, ensure_ascii=False, separators=(',', ':'))
pos_tags_js = json.dumps(pos_tags,   ensure_ascii=False, separators=(',', ':'))
katakana_js = json.dumps(katakana,   ensure_ascii=False, separators=(',', ':'))
same_meaning_js = json.dumps(same_meaning, ensure_ascii=False, separators=(',', ':'))

html = replace_const(html, 'KANJI',    kanji_js)
html = replace_const(html, r'SKGS ',   skgs_js)   # 'SKGS ' matches 'const SKGS  ='
html = replace_const(html, 'VOCAB',    vocab_js)
html = replace_const(html, 'POS_TAGS', pos_tags_js)
html = replace_const(html, 'KATAKANA', katakana_js)
html = replace_const(html, 'SAME_MEANING', same_meaning_js)

tz      = ZoneInfo('Europe/Warsaw')
now_str = datetime.now(tz).strftime('%Y-%m-%d %H:%M')
html    = re.sub(r'updated \d{4}-\d{2}-\d{2} \d{2}:\d{2}', f'updated {now_str}', html)

with open(INDEX_HTML, 'w', encoding='utf-8') as f:
    f.write(html)
print(f"  Timestamp → {now_str}")

# ── 6. Bump SW cache version ──────────────────────────────────────────────────
print("Bumping SW cache version...")
with open(SW_JS, 'r', encoding='utf-8') as f:
    sw = f.read()

m = re.search(r"const CACHE = 'koku-v(\d+)'", sw)
if m:
    old_v = int(m.group(1))
    new_v = old_v + 1
    sw    = sw.replace(f"'koku-v{old_v}'", f"'koku-v{new_v}'")
    with open(SW_JS, 'w', encoding='utf-8') as f:
        f.write(sw)
    print(f"  koku-v{old_v} → koku-v{new_v}")
else:
    print("  WARNING: could not find cache version in sw.js")

# ── Done ──────────────────────────────────────────────────────────────────────
print()
print("✓ All done! Upload index.html + sw.js to GitHub.")
input("Press Enter to close...")
