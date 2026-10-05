"""
Same-Meaning Detector
─────────────────────────────────────────────────────────────────────
Finds same-meaning / different-kanji word groups in your CURRENT vocab
that are NOT yet covered by FINAL SAME MEANING LIST, and writes them to
'same_meaning_NEW_candidates.xlsx' for annotation.

Workflow:
  1. Add your new vocab to FINAL VOCAB LIST (and run update_koku.bat as usual).
  2. Run this script (double-click or `python same_meaning_detector.py`).
  3. Send 'same_meaning_NEW_candidates.xlsx' to Claude to get the nuance notes.
  4. Paste the finished rows into FINAL SAME MEANING LIST, bump its version, run the .bat.

Read-only on your existing spreadsheets — it never edits them.
"""

import openpyxl, glob, os, re
from collections import defaultdict

APP_DIR = os.path.dirname(os.path.abspath(__file__))


def latest(pattern):
    m = glob.glob(os.path.join(APP_DIR, pattern))
    return max(m, key=os.path.getmtime) if m else None


def normalise_meaning(m):
    m = m.lower().strip()
    parts = re.split(r'[|;/,]', m)
    return [re.sub(r'^(to |a |an |the )', '', x.strip()).strip() for x in parts if x.strip()]


# ── 1. current vocab → candidate groups ───────────────────────────────────────
vf = latest('FINAL VOCAB LIST*.xlsx')
if not vf:
    print('ERROR: FINAL VOCAB LIST*.xlsx not found in this folder.')
    input('Press Enter to exit...'); raise SystemExit
print('Reading (read-only):', os.path.basename(vf))

wb = openpyxl.load_workbook(vf, read_only=True, data_only=True)
ws = wb['Sheet1'] if 'Sheet1' in wb.sheetnames else wb.active

seen = set(); items = []
for row in ws.iter_rows(values_only=True):
    if not row or not row[0]:
        continue
    for c in range(5, len(row)):
        cell = row[c]
        if not cell:
            continue
        p = str(cell).strip().split('\n')
        if len(p) >= 3:
            w = p[0].strip()
            if w and w not in seen:
                seen.add(w)
                items.append((w, p[1].strip(), p[2].strip()))

groups = defaultdict(dict)   # normalised meaning -> {word: reading}
for w, r, m in items:
    for nm in normalise_meaning(m):
        groups[nm][w] = r
candidates = {nm: list(d.items()) for nm, d in groups.items() if len(d) >= 2}

# ── 2. existing same-meaning list → words already covered ──────────────────────
existing_words = set()
smf = latest('FINAL SAME MEANING LIST*.xlsx')
if smf:
    print('Comparing against:', os.path.basename(smf))
    wbs = openpyxl.load_workbook(smf, read_only=True, data_only=True)
    wss = wbs[wbs.sheetnames[0]]
    for row in wss.iter_rows(min_row=2, values_only=True):
        for i in (1, 3, 5, 7, 9, 11):   # the six Word columns (cols 2-13)
            if i < len(row) and row[i]:
                existing_words.add(str(row[i]).strip())
else:
    print('No FINAL SAME MEANING LIST found — every candidate will be treated as new.')

# Meanings deliberately rejected during curation (English homonyms / non-synonyms).
# These would otherwise be re-flagged every run because their words are never added
# to the list. Add a meaning here (lowercase, no leading a/an/the) to silence it.
REJECTED = {
    'ad', 'be raised', 'character', 'close', 'compromise', 'conclusion', 'face',
    'kind', 'mercury', 'nature', 'order', 'ring', 'second', 'spring', 'state',
    'superior', 'typical',
    # 'support' curated as 応援/支援 (row 29); 支援する intentionally excluded (same kanji as 支援)
    'support',
    'tie',   # homonym: 結ぶ (tie a knot) vs 引き分け (a drawn match) — not synonyms
    'fine',  # homonym: 結構 (fine = okay/good) vs 罰金 (a fine = monetary penalty)
    'vending machine',  # 自販機 is just the contraction of 自動販売機 — same word, not a nuance pair
    'export',           # 輸出する is just the verb form of 輸出 — same kanji, not a nuance pair
}

# ── 3. new candidates = groups with a not-yet-covered word, minus rejected ─────
new = []
for nm, ws_ in sorted(candidates.items()):
    if nm in REJECTED:
        continue
    if any(w not in existing_words for w, r in ws_):
        new.append((nm, ws_))

print(f'\nVocab words scanned: {len(items)}')
print(f'Candidate groups: {len(candidates)} | NEW (need notes): {len(new)}')

out = openpyxl.Workbook(); o = out.active; o.title = 'New Candidates'
o.append(['Meaning', 'Word 1', 'Reading 1', 'Word 2', 'Reading 2',
          'Word 3', 'Reading 3', 'Word 4', 'Reading 4',
          'Word 5', 'Reading 5', 'Word 6', 'Reading 6', 'Note'])
for nm, ws_ in new:
    rowv = [nm.title()]
    for i in range(6):
        rowv += [ws_[i][0], ws_[i][1]] if i < len(ws_) else ['', '']
    rowv.append('')
    o.append(rowv)
outpath = os.path.join(APP_DIR, 'same_meaning_NEW_candidates.xlsx')
out.save(outpath)
print(f'\nSaved: {os.path.basename(outpath)}')
print('→ Send that file to Claude for the nuance notes, then paste the rows into FINAL SAME MEANING LIST.')
input('Press Enter to exit...')
